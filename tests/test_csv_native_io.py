"""Focused native import parity, package preservation and transaction checks."""

import ast
import csv
import importlib.machinery
import importlib.util
import io
import os
import random
import shutil
import types
import typing
import zipfile
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

import gl260_csv_io as writer


@pytest.fixture(scope="module")
def native():
    """Load the requested built extension, or installed backend, for native tests.

    No inputs; returns a module or skips if kernels are absent. Loads a DLL but
    does not install it or modify the running application; import errors propagate.
    """
    path = os.environ.get("GL260_TEST_RUST_PATH")
    if path:
        loader = importlib.machinery.ExtensionFileLoader("gl260_rust_ext", path)
        spec = importlib.util.spec_from_file_location(
            "gl260_rust_ext", path, loader=loader
        )
        module = importlib.util.module_from_spec(spec)
        loader.exec_module(module)
    else:
        module = pytest.importorskip("gl260_rust_ext")
    if not hasattr(module, "write_csv_workbook_core"):
        pytest.skip("Installed extension predates CSV import kernels")
    return module


@pytest.fixture(scope="module")
def app():
    """Extract actual import functions without launching Tk or changing settings.

    Returns their execution namespace with backend controls stubbed for tests.
    Reads the entry file; parsing errors propagate and no production files change.
    """
    import re

    source = Path("GL-260 Data Analysis and Plotter.py").read_text(encoding="utf-8-sig")
    names = {
        "_parse_gl260_csv_table_with_metadata",
        "_extract_gl260_csv_table_with_metadata",
        "_extract_graphtec_channel_metadata_row",
        "_gl260_row_has_values",
        "_gl260_row_is_data_marker",
        "_normalize_gl260_header",
        "_normalize_csv_channel_metadata",
        "_csv_channel_display_label",
        "_write_gl260_output_sheet",
        "_write_gl260_output_sheet_python",
        "_sanitize_excel_sheet_name",
        "_csv_import_native_backend",
    }
    selected = [
        n
        for n in ast.parse(source).body
        if isinstance(n, ast.FunctionDef) and n.name in names
    ]
    namespace = dict(
        globals(),
        **vars(typing),
        re=re,
        CSV_IMPORT_DEFAULT_SHEET_EXISTS_MODE="error",
        CSV_IMPORT_SHEET_EXISTS_MODES=("error", "overwrite", "autosuffix"),
    )
    namespace["_current_rust_backend_mode"] = lambda: "auto"
    namespace["_is_rust_kernel_session_unhealthy"] = lambda name: False
    namespace["_mark_rust_kernel_session_unhealthy"] = lambda *args, **kwargs: None
    namespace["_load_rust_backend"] = lambda: None
    exec(
        compile(ast.Module(body=selected, type_ignores=[]), "<csv-production>", "exec"),
        namespace,
    )
    return namespace


def _fixture_book(path, epoch_1904=False):
    """Write a small source workbook with formulas, styles and an opaque ZIP part.

    Inputs are a temporary path and epoch flag. Returns None; writes test fixtures
    only, allowing byte-level preservation checks. I/O errors propagate.
    """
    book = openpyxl.Workbook()
    book.active.title = "Original"
    book.active["A1"] = "Keep me"
    book.active["B2"] = "=1+2"
    book.active["A1"].font = openpyxl.styles.Font(bold=True, color="FF123456")
    book.active.merge_cells("C3:D4")
    if epoch_1904:
        book.epoch = openpyxl.utils.datetime.CALENDAR_MAC_1904
    book.save(path)
    book.close()
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("customXml/item1.xml", b"<custom>preserve exactly</custom>")


def _frame():
    """Return representative date/numeric/text import data without side effects."""
    return pd.DataFrame(
        {
            "Date & Time": pd.to_datetime(
                ["2026-10-30 12:00:00.123", "2026-10-30 12:00:01", None],
                format="ISO8601",
            ),
            "Pressure": [1.25, np.nan, -4.5],
            "Raw": [" a&<b>\r\n ", "=SUM(B2:B3)", "#N/A"],
            "Boolean": [True, False, True],
        }
    )


def _no_fallback(*args):
    """Raise if a native success test unexpectedly invokes Python; no mutations."""
    raise AssertionError("Unexpected Python fallback")


@pytest.mark.parametrize("epoch_1904", [False, True])
def test_native_preserves_package_and_values(tmp_path, native, app, epoch_1904):
    """Write temporary native/Python results; compare cells, styles and ZIP parts.

    Fixture inputs supply native and production functions. Returns None; assertions
    fail on parity/preservation loss, and only pytest temporary files are changed.
    """
    target = tmp_path / "native.xlsx"
    _fixture_book(target, epoch_1904)
    with zipfile.ZipFile(target) as archive:
        before = {name: archive.read(name) for name in archive.namelist()}
    reference = tmp_path / "python.xlsx"
    shutil.copy2(target, reference)
    frame = _frame()
    assert (
        writer.write_import_workbook(
            target, "Imported", frame, "error", native, _no_fallback
        )
        == "Imported"
    )
    app["_write_gl260_output_sheet_python"](reference, "Imported", frame, "error")
    with zipfile.ZipFile(target) as archive:
        for name, content in before.items():
            if name not in {
                "xl/workbook.xml",
                "xl/_rels/workbook.xml.rels",
                "[Content_Types].xml",
                "xl/styles.xml",
            }:
                assert archive.read(name) == content
    actual = openpyxl.load_workbook(target)
    expected = openpyxl.load_workbook(reference)
    try:
        assert list(actual["Imported"].values) == list(expected["Imported"].values)
        assert actual["Imported"].freeze_panes == "A2"
        assert actual["Original"]["A1"].font.bold
        assert actual["Original"]["B2"].value == "=1+2"
        assert str(actual["Original"].merged_cells) == "C3:D4"
        assert (
            actual["Imported"]["A2"].number_format
            == expected["Imported"]["A2"].number_format
        )
    finally:
        actual.close()
        expected.close()


@pytest.mark.parametrize("mode", ["overwrite", "autosuffix", "error"])
def test_collision_modes(tmp_path, native, app, mode):
    """Verify collision semantics and unchanged original on errors using fixtures.

    Returns None; writes temporary workbooks and raises on assertion failure.
    """
    target = tmp_path / "collision.xlsx"
    _fixture_book(target)
    before = target.read_bytes()
    if mode == "error":
        with pytest.raises(ValueError, match="already exists"):
            writer.write_import_workbook(
                target,
                "Original",
                _frame(),
                mode,
                native,
                app["_write_gl260_output_sheet_python"],
            )
        assert target.read_bytes() == before
    else:
        name = writer.write_import_workbook(
            target, "Original", _frame(), mode, native, _no_fallback
        )
        assert name == ("Original_1" if mode == "autosuffix" else "Original")
    assert not list(tmp_path.glob(".gl260-import-*"))


def test_python_fallback_and_failure_atomicity(tmp_path, app, monkeypatch):
    """Force native failure and validation rejection; prove cleanup and atomicity.

    Fixtures provide temporary paths and patch restoration. Returns None; only
    temporary workbooks and in-test callbacks change. Assertions expose failures.
    """
    target = tmp_path / "fallback.xlsx"
    _fixture_book(target)
    failures = []

    def broken(*args):
        """Simulate a failed native write; accepts kernel args and always raises."""
        Path(args[1]).write_bytes(b"partial ZIP")
        raise RuntimeError("injected native failure")

    backend = types.SimpleNamespace(write_csv_workbook_core=broken)
    writer.write_import_workbook(
        target,
        "Imported",
        _frame(),
        "error",
        backend,
        app["_write_gl260_output_sheet_python"],
        lambda *a, **k: failures.append((a, k)),
    )
    assert failures
    before = target.read_bytes()

    def reject(*args, **kwargs):
        """Reject any candidate workbook to verify no commit; always raises."""
        raise ValueError("injected validation failure")

    monkeypatch.setattr(writer, "_validate_output", reject)
    with pytest.raises(ValueError, match="injected validation"):
        writer.write_import_workbook(
            target,
            "Other",
            _frame(),
            "error",
            None,
            app["_write_gl260_output_sheet_python"],
        )
    assert target.read_bytes() == before
    assert not list(tmp_path.glob(".gl260-import-*"))


@pytest.mark.parametrize("encoding", ["utf-8-sig", "cp1252", "latin-1"])
def test_csv_encoding_and_preview_parity(tmp_path, native, app, monkeypatch, encoding):
    """Compare production native/fallback parsing and streaming preview per encoding.

    Writes only fixture CSV data; monkeypatch restores the backend. Returns None;
    assertions include metadata, duplicate headers and irregular row widths.
    """
    path = tmp_path / "input.csv"
    text = 'CH,Signal Name,Unit\r\nCH1,Temp é,°C\r\nData\r\nDate&Time,CH1,CH1,\r\ns,C,C,\r\n2026-01-01,+2," a,b ","quoted\r\nline"\r\n2026-01-02,-,4\r\n2026-01-03,5,6,7,extra\r\n'
    if encoding == "latin-1":
        text = text.replace("Temp é", "Temp \x81")
    path.write_bytes(text.encode(encoding))
    parse = app["_parse_gl260_csv_table_with_metadata"]
    monkeypatch.setitem(app, "_load_rust_backend", lambda: None)
    expected = parse(path)
    monkeypatch.setitem(app, "_load_rust_backend", lambda: native)
    assert parse(path) == expected
    assert parse(path, max_rows=1) == (expected[0], expected[1][:1], expected[2])


def test_csv_randomized_tokenization(native, app):
    """Compare native parsing with Python csv.reader on generated quoted records.

    Uses fixed random data including multiline text; returns None, changes no
    files, and raises on row/metadata differences.
    """
    rng = random.Random(260)
    stream = io.StringIO(newline="")
    csv_writer = csv.writer(stream)
    csv_writer.writerows([["Data"], ["Date&Time", "CH1", "CH2"], ["s", "PSI", "PSI"]])
    alphabet = 'abc, "\r\n\t+é\x1c'
    for _ in range(1000):
        csv_writer.writerow(
            ["2026-01-01"]
            + [
                "".join(rng.choice(alphabet) for _ in range(rng.randrange(20)))
                for _ in range(rng.randrange(5))
            ]
        )
    text = stream.getvalue()
    extract = app["_extract_gl260_csv_table_with_metadata"]
    expected = extract(csv.reader(io.StringIO(text, newline="")))
    prefix, rows = native.csv_table_core(text, csv.field_size_limit())
    columns, _, metadata = extract(prefix)
    assert (columns, rows, metadata) == expected


@pytest.mark.parametrize(
    "xml",
    [
        '<worksheet><sheetData><row r="1"/></sheetData></worksheet>',
        '<worksheet><sheetData><row r="1"/><row r="3"/></sheetData></worksheet>',
        '<worksheet><sheetData><row r="1"/><row r="2"/></sheetData>',
    ],
)
def test_native_validator_rejects_incomplete_xml(tmp_path, native, xml):
    """Reject truncated/missing/out-of-order rows in a temporary candidate ZIP.

    Inputs supply native backend and corrupt XML. Returns None; writes test data
    only and asserts the independent native reader raises before commit.
    """
    path = tmp_path / "invalid.xlsx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("xl/worksheets/sheet1.xml", xml)
    with pytest.raises(RuntimeError):
        native.validate_csv_worksheet_core(str(path), "xl/worksheets/sheet1.xml", 1)


def test_replace_failure_keeps_original(tmp_path, app, monkeypatch):
    """Inject a locked-destination failure after validation and verify zero commit.

    Uses temporary fixtures; returns None. Patches os.replace for this test only,
    proving the original bytes survive and the candidate is removed on failure.
    """
    target = tmp_path / "locked.xlsx"
    _fixture_book(target)
    before = target.read_bytes()

    def locked(*args):
        """Simulate an Excel-held destination; always raise without modifying files."""
        raise PermissionError("destination locked")

    monkeypatch.setattr(writer.os, "replace", locked)
    with pytest.raises(PermissionError):
        writer.write_import_workbook(
            target,
            "Imported",
            _frame(),
            "error",
            None,
            app["_write_gl260_output_sheet_python"],
        )
    assert target.read_bytes() == before
    assert not list(tmp_path.glob(".gl260-import-*"))


def test_parser_failure_and_python_mode(tmp_path, app, monkeypatch):
    """Verify failed native parsing retries Python and Python-only mode bypasses it.

    Uses a temporary CSV and restored backend stubs. Returns None; raises on
    incorrect routing, data loss or missing health diagnostics.
    """
    path = tmp_path / "fallback.csv"
    path.write_text("Data\nDate&Time,CH1\ns,PSI\n2026-01-01,2\n", encoding="utf-8")
    calls = []

    def broken(*args):
        """Record a native attempt and raise to exercise parser fallback."""
        calls.append("native")
        raise RuntimeError("injected parser failure")

    backend = types.SimpleNamespace(csv_table_core=broken)
    monkeypatch.setitem(app, "_load_rust_backend", lambda: backend)
    monkeypatch.setitem(
        app,
        "_mark_rust_kernel_session_unhealthy",
        lambda *a, **k: calls.append("unhealthy"),
    )
    result = app["_parse_gl260_csv_table_with_metadata"](path)
    assert result[1] == [["2026-01-01", "2"]]
    assert calls == ["native", "unhealthy"]
    monkeypatch.setitem(app, "_current_rust_backend_mode", lambda: "python")
    assert app["_parse_gl260_csv_table_with_metadata"](path) == result
    assert calls == ["native", "unhealthy"]
