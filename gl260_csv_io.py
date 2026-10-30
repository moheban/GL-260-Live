"""Transactional CSV-import workbook output with optional Rust acceleration.

Only temporary sibling files are written before validation. Native output copies
unrelated package entries verbatim; unsupported workbooks use the legacy writer.
"""

import os
import posixpath
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from xml.dom import minidom

import numpy as np
import pandas as pd

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"


class UnsupportedWorkbook(ValueError):
    """Signal a known native limitation that should use Python without health loss."""


def _child(parent, namespace, name, attributes):
    """Append and return a namespaced DOM element with string attributes.

    Mutates the supplied metadata DOM to preserve namespace declarations and
    extension prefixes. DOM errors propagate; no files are touched.
    """
    prefix = parent.prefix
    node = parent.ownerDocument.createElementNS(
        namespace, f"{prefix}:{name}" if prefix else name
    )
    for key, value in attributes.items():
        node.setAttribute(key, str(value))
    parent.appendChild(node)
    return node


def _package_plan(archive, name, mode):
    """Plan minimal workbook metadata edits, preserving unrelated ZIP members.

    Inputs are an open ZipFile, sanitized sheet title and collision mode. Returns
    resolved title, worksheet part, replacement XML bytes, date style and epoch.
    Mutates only in-memory DOMs. Unsupported package features raise ValueError
    so the caller can use the established Python writer.
    """
    names = archive.namelist()
    if len(names) != len(set(names)) or any(
        n.startswith("_xmlsignatures/") or n == "xl/calcChain.xml" for n in names
    ):
        raise UnsupportedWorkbook(
            "Native writer does not support signed/ambiguous packages."
        )
    workbook = minidom.parseString(archive.read("xl/workbook.xml"))
    relations = minidom.parseString(archive.read("xl/_rels/workbook.xml.rels"))
    types = minidom.parseString(archive.read("[Content_Types].xml"))
    styles = minidom.parseString(archive.read("xl/styles.xml"))
    sheets = workbook.getElementsByTagNameNS(MAIN, "sheets")[0]
    entries = list(sheets.getElementsByTagNameNS(MAIN, "sheet"))
    existing = {s.getAttribute("name") for s in entries}
    original_name = name
    if name in existing:
        if mode == "error":
            raise UnsupportedWorkbook(f"Sheet '{name}' already exists in the workbook.")
        if mode == "autosuffix":
            counter = 1
            while name in existing:
                suffix = f"_{counter}"
                name = original_name[: 31 - len(suffix)] + suffix
                counter += 1
    # openpyxl handles case-insensitive title collisions and complex overwrite
    # references; defer these uncommon cases rather than guessing Excel rules.
    if name not in existing and name.casefold() in {s.casefold() for s in existing}:
        raise UnsupportedWorkbook(
            "Case-insensitive sheet collision requires Python writer."
        )
    rels = list(relations.getElementsByTagNameNS(PKG, "Relationship"))
    overwrite = next((s for s in entries if s.getAttribute("name") == name), None)
    if overwrite is not None:
        rid = overwrite.getAttributeNS(REL, "id")
        relation = next(r for r in rels if r.getAttribute("Id") == rid)
        target = relation.getAttribute("Target")
        part = posixpath.normpath(
            target.lstrip("/") if target.startswith("/") else "xl/" + target
        )
        sheet_rels = posixpath.join(
            posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels"
        )
        if sheet_rels in names or workbook.getElementsByTagNameNS(MAIN, "definedName"):
            raise UnsupportedWorkbook(
                "Referenced worksheet overwrite requires Python writer."
            )
        if not part.startswith("xl/worksheets/") or part not in names:
            raise UnsupportedWorkbook("Unsupported worksheet relationship.")
        sheets.removeChild(overwrite)
        overwrite.setAttribute("state", "visible")
        sheets.appendChild(overwrite)
    else:
        number = 1
        while f"xl/worksheets/sheet{number}.xml" in names:
            number += 1
        part = f"xl/worksheets/sheet{number}.xml"
        relation_number = 1
        used = {r.getAttribute("Id") for r in rels}
        while f"rId{relation_number}" in used:
            relation_number += 1
        rid = f"rId{relation_number}"
        sheet = _child(
            sheets,
            MAIN,
            "sheet",
            {
                "name": name,
                "sheetId": max(int(s.getAttribute("sheetId")) for s in entries) + 1,
            },
        )
        sheet.setAttributeNS("http://www.w3.org/2000/xmlns/", "xmlns:r", REL)
        sheet.setAttributeNS(REL, "r:id", rid)
        _child(
            relations.documentElement,
            PKG,
            "Relationship",
            {"Id": rid, "Type": REL + "/worksheet", "Target": "/" + part},
        )
        _child(
            types.documentElement,
            CT,
            "Override",
            {
                "PartName": "/" + part,
                "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml",
            },
        )
    # Append a date style instead of renumbering existing style references.
    xfs = styles.getElementsByTagNameNS(MAIN, "cellXfs")[0]
    date_style = len(xfs.getElementsByTagNameNS(MAIN, "xf"))
    formats = styles.getElementsByTagNameNS(MAIN, "numFmts")
    if not formats:
        raise UnsupportedWorkbook("Missing number-format table requires Python writer.")
    formats = formats[0]
    format_id = (
        max(
            [163]
            + [
                int(f.getAttribute("numFmtId"))
                for f in formats.getElementsByTagNameNS(MAIN, "numFmt")
            ]
        )
        + 1
    )
    _child(
        formats,
        MAIN,
        "numFmt",
        {"numFmtId": format_id, "formatCode": "yyyy-mm-dd h:mm:ss"},
    )
    formats.setAttribute(
        "count", str(len(formats.getElementsByTagNameNS(MAIN, "numFmt")))
    )
    _child(
        xfs,
        MAIN,
        "xf",
        {"numFmtId": format_id, "fontId": 0, "fillId": 0, "borderId": 0, "xfId": 0},
    )
    xfs.setAttribute("count", str(date_style + 1))
    props = workbook.getElementsByTagNameNS(MAIN, "workbookPr")
    epoch_1904 = bool(props and props[0].getAttribute("date1904") in ("1", "true"))
    replacements = {
        "xl/workbook.xml": workbook.toxml(encoding="utf-8"),
        "xl/_rels/workbook.xml.rels": relations.toxml(encoding="utf-8"),
        "[Content_Types].xml": types.toxml(encoding="utf-8"),
        "xl/styles.xml": styles.toxml(encoding="utf-8"),
    }
    return name, part, replacements, date_style, epoch_1904


def _column_buffers(frame, epoch_1904):
    """Return native numeric/text buffers, kinds and slot indexes for a DataFrame.

    Converts naive dates to workbook-epoch days and missing values to empty cells.
    Inputs are the import frame and epoch flag. No files or input data are changed.
    Unsupported object types/timezones raise ValueError for Python fallback.
    """
    from openpyxl.xml import LXML

    numeric, text, kinds, slots = [], [], [], []
    for _, series in frame.items():
        if pd.api.types.is_datetime64_any_dtype(series.dtype):
            if getattr(series.dt, "tz", None) is not None:
                raise UnsupportedWorkbook("Timezone-aware dates require Python writer.")
            origin = pd.Timestamp("1904-01-01" if epoch_1904 else "1899-12-30")
            values = (series - origin).dt.total_seconds().to_numpy(dtype=float) / 86400
            if not epoch_1904:
                values[(values >= 1) & (values < 61)] -= 1
            kind = "date"
        elif pd.api.types.is_numeric_dtype(series.dtype):
            values = series.to_numpy(dtype=float, na_value=np.nan)
            kind = "bool" if pd.api.types.is_bool_dtype(series.dtype) else "number"
        else:
            values = []
            for value in series:
                if value is None or value is pd.NA or value is pd.NaT:
                    values.append(None)
                elif isinstance(value, str):
                    # openpyxl's etree serializer normalizes CR to LF, whereas
                    # lxml preserves CR using a character reference. Match the
                    # installed fallback so switching backend retains cell text.
                    values.append(value if LXML else value.replace("\r", "\n"))
                else:
                    raise UnsupportedWorkbook(
                        "Mixed object column requires Python writer."
                    )
            slots.append(len(text))
            kinds.append("text")
            text.append(values)
            continue
        slots.append(len(numeric))
        kinds.append(kind)
        numeric.append(values)
    matrix = np.ascontiguousarray(numeric, dtype=float)
    if not numeric:
        matrix = np.empty((0, len(frame)), dtype=float)
    return matrix, text, kinds, slots


def _validate_output(path, title, frame, original=None, changed=(), backend=None):
    """Validate a temporary XLSX before committing it to the user's destination.

    Inputs identify the temporary path, expected sheet/frame and optional original
    ZIP plus changed parts and optional native XML validator. Checks CRCs,
    unchanged member identities, worksheet
    XML, row count and readable headers. Returns None; reads files only and raises
    ValueError/ZIP/XML/openpyxl errors on corruption or schema mismatch.
    """
    import openpyxl

    with zipfile.ZipFile(path) as archive:
        if len(archive.namelist()) != len(set(archive.namelist())) or archive.testzip():
            raise ValueError("Workbook ZIP validation failed.")
        if original is not None:
            for info in original.infolist():
                if info.filename not in changed:
                    copied = archive.getinfo(info.filename)
                    if (copied.CRC, copied.file_size) != (info.CRC, info.file_size):
                        raise ValueError(
                            f"Unrelated workbook part changed: {info.filename}"
                        )
        for name in changed:
            if not name.startswith("xl/worksheets/") and (
                name.endswith(".xml") or name.endswith(".rels")
            ):
                with archive.open(name) as handle:
                    # Clear each parsed node to keep validation memory bounded.
                    for _, element in ET.iterparse(handle, events=("end",)):
                        element.clear()
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = book[title]
        if sheet.max_row != len(frame) + 1 or sheet.max_column != len(frame.columns):
            raise ValueError("Imported worksheet dimensions did not match the input.")
        headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        if list(headers) != list(frame.columns):
            raise ValueError("Imported worksheet headers did not match the input.")
        validate = getattr(backend, "validate_csv_worksheet_core", None)
        if callable(validate):
            validate(str(path), sheet._worksheet_path, len(frame))
            return
        # Inspect actual rows as well as the declared dimension. This catches a
        # truncated-but-well-formed worksheet before either writer can commit it.
        with zipfile.ZipFile(path) as archive:
            with archive.open(sheet._worksheet_path) as handle:
                expected_row = 1
                for _, element in ET.iterparse(handle, events=("end",)):
                    if element.tag == f"{{{MAIN}}}row":
                        if int(element.get("r", "0")) != expected_row:
                            raise ValueError(
                                "Imported worksheet row sequence is incomplete."
                            )
                        expected_row += 1
                    element.clear()
                if expected_row != len(frame) + 2:
                    raise ValueError("Imported worksheet data rows are incomplete.")
    finally:
        book.close()


def write_import_workbook(
    path, title, frame, mode, backend, python_writer, on_native_error=None
):
    """Write and validate a sibling workbook, then atomically replace destination.

    Inputs are destination path, sanitized title, DataFrame, collision mode,
    optional native backend, legacy writer and optional health callback callable.
    Returns resolved title; callback receives failed kernel name/reason/detail.
    Native failures retry through Python; both paths validate before replacement.
    Temporary files are always deleted. I/O/validation failures propagate while
    leaving the destination untouched. Concurrent destination changes abort.
    """
    import shutil

    destination = Path(path)
    before = destination.stat()
    descriptor, temporary = tempfile.mkstemp(
        prefix=".gl260-import-", suffix=destination.suffix, dir=destination.parent
    )
    os.close(descriptor)
    try:
        native = getattr(backend, "write_csv_workbook_core", None)
        used_native = False
        if callable(native):
            try:
                with zipfile.ZipFile(destination) as original:
                    resolved, part, patches, style, epoch = _package_plan(
                        original, title, mode
                    )
                    buffers = _column_buffers(frame, epoch)
                    native(
                        str(destination),
                        temporary,
                        part,
                        patches,
                        list(frame.columns),
                        *buffers,
                        style,
                        len(frame),
                    )
                    _validate_output(
                        temporary, resolved, frame, original, {*patches, part}, backend
                    )
                used_native = True
            except UnsupportedWorkbook:
                used_native = False
            except Exception as exc:
                # No timed-out background writer can race this retry: the native
                # call is synchronous and touches only this request's temporary.
                used_native = False
                if on_native_error is not None:
                    on_native_error(
                        "write_csv_workbook_core",
                        reason="invalid_payload",
                        detail=str(exc),
                    )
        if not used_native:
            shutil.copy2(destination, temporary)
            resolved = python_writer(temporary, title, frame, mode)
            _validate_output(temporary, resolved, frame)
        after = destination.stat()
        if (before.st_mtime_ns, before.st_size, before.st_ino) != (
            after.st_mtime_ns,
            after.st_size,
            after.st_ino,
        ):
            raise RuntimeError(
                "Destination workbook changed during import; retry on the latest file."
            )
        # Windows FlushFileBuffers requires a writable handle, even after all
        # serializer streams have closed and the candidate has been validated.
        with open(temporary, "r+b") as handle:
            os.fsync(handle.fileno())
        os.replace(temporary, destination)
        return resolved
    finally:
        Path(temporary).unlink(missing_ok=True)
